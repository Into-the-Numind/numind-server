"""
template.py — Template loading and variable substitution for xlsx-author skill.

Templates are pre-built .xlsx files stored at
``/skills/xlsx-author/templates/``.  Placeholders in cells use the
``{{variable_name}}`` syntax and are replaced at load time.

Available templates:
    summary       -- Three-sheet: Cover + Data + Charts
    daily-report  -- Single-sheet daily/weekly report with KPI section
    comparison    -- Side-by-side comparison with conditional formatting
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.workbook.workbook import Workbook

# Template directory mounted inside the sandbox container.
TEMPLATE_DIR = Path("/skills/xlsx-author/templates")

# Matches {{variable_name}} placeholders (alphanumeric + underscore).
_PLACEHOLDER_RE = re.compile(r"\{\{(\w+)\}\}")


def load_template(template_name: str, vars: dict[str, str]) -> Workbook:
    """Load a pre-built .xlsx template and substitute ``{{key}}`` placeholders.

    Iterates over every cell in every worksheet and performs regex substitution
    for each ``{{key}}`` found using values from *vars*.  Placeholders whose
    key is not in *vars* are left unchanged.

    Args:
        template_name:  One of "summary", "daily-report", "comparison".
        vars:           Mapping of placeholder name → replacement string.
                        Values should be strings; non-strings are converted with
                        ``str()``.

    Returns:
        An openpyxl ``Workbook`` object with all placeholders replaced.
        Callers can add more sheets or data, then save the workbook.

    Raises:
        FileNotFoundError: When the template file does not exist.
        ValueError:        When ``template_name`` contains path traversal
                           characters (security check).

    Example::

        wb = load_template("summary", {
            "title": "2026年Q2销售汇总",
            "period": "2026-04 ~ 2026-06",
            "generated_date": "2026-07-01",
            "summary_text": "本季度营收同比增长 23%",
        })
        wb.save("/output/Q2报告.xlsx")
    """
    # Security: reject path traversal attempts.
    if ".." in template_name or "/" in template_name or "\\" in template_name:
        raise ValueError(f"Invalid template name: {template_name!r}")

    template_path = TEMPLATE_DIR / f"{template_name}.xlsx"
    if not template_path.exists():
        raise FileNotFoundError(
            f"Template not found: {template_path}. "
            f"Available templates are in {TEMPLATE_DIR}."
        )

    wb = openpyxl.load_workbook(str(template_path))

    # String values for substitution.
    str_vars = {k: str(v) for k, v in vars.items()}

    def _replace(m: re.Match) -> str:  # type: ignore[type-arg]
        key = m.group(1)
        return str_vars.get(key, m.group(0))

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = _PLACEHOLDER_RE.sub(_replace, cell.value)

    return wb
